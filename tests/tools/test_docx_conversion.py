"""Actual, deterministic DOCX parsing; no inference models or external APIs needed."""

from __future__ import annotations

import hashlib

import pytest

pytest.importorskip("docling")

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from src.core.uploads import UploadRecord
from src.infra.docling_adapter import convert_file
from src.infra.document_ingestion import ConversionPolicy


def convert_word(word, tmp_path):
    path = tmp_path / "formatted.docx"
    word.save(path)
    raw = path.read_bytes()
    file = UploadRecord(file_id="formatted", name=path.name, path=str(path), size_bytes=len(raw),
                        content_hash="sha256:" + hashlib.sha256(raw).hexdigest(),
                        source_uri="upload://local/formatted.docx")
    # The Word pipeline is declarative and does not load the PDF/OCR models.
    return convert_file(file, ConversionPolicy(artifacts_path=str(tmp_path)))


def test_code_keeps_indentation_and_blank_lines_separate_from_body(tmp_path):
    """Code formatting preserves its block and indentation without absorbing the following paragraph."""
    word = Document()
    word.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    word.add_heading("Usage", 1)
    word.add_paragraph("if flag:", style="Code")
    paragraph = word.add_paragraph("    use(", style="Code")
    paragraph.add_run("24").bold = True
    paragraph.add_run(")")
    word.add_paragraph("", style="Code")
    word.add_paragraph("\tfinish()", style="Code")
    word.add_paragraph("Following body paragraph.")
    parsed = convert_word(word, tmp_path)
    assert [(e.kind, e.text, e.heading_path) for e in parsed.elements] == [
        ("heading", "Usage", []),
        ("code", "if flag:\n    use(24)\n\n\tfinish()", ["Usage"]),
        ("paragraph", "Following body paragraph.", ["Usage"]),
    ]
