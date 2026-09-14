"""Build small, original bilingual fixtures for document ingestion tests.

Run with the bundled artifact Python via ``uv run --no-project --python ...``.
The source text is deliberately owned by this repository; no customer documents
or copyrighted example papers are downloaded for the OCR benchmark.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import shutil
import subprocess


TITLE = "문서 처리 실험 Document processing study"
HEADING = "학습 자료 Learning notes"
LINES = [
    "문서 검색은 원본 페이지와 표의 구조를 함께 보존합니다.",
    "Document retrieval preserves source pages and table structure.",
    "실험 식별자는 DOCUMATE-2026이며 표본 수는 24개입니다.",
    "The experiment identifier is DOCUMATE-2026 and the sample size is 24.",
    "한국어와 English text를 같은 문단에서 읽고 검증합니다.",
]
TABLE = [
    ["항목 Item", "수량 Count", "결과 Result"],
    ["학습 Training", "16", "완료 Complete"],
    ["평가 Evaluation", "8", "대기 Pending"],
]
PAGE_TWO = [
    "추가 검증 Additional validation",
    "두 번째 페이지의 정답은 42입니다.",
    "The answer on the second page is 42.",
]

# Exact logical text, independent of the run boundaries used below. Local
# conversion tests import these values to check whitespace and source spans.
MIXED_TITLE = "Mixed formatting document"
MIXED_HEADING = "Formatting boundary checks"
MIXED_INLINE = "The silver fox follows the amber trail safely."
MIXED_TOKEN = "The code is DOCUMATE-2026 and parse_document."
MIXED_KOREAN = "한글서식경계를그대로보존합니다."
MIXED_SPACING = "Before  two spaces\tTabbed segment\nSoft line break stays in this paragraph."
MIXED_ADJACENT = (
    "First adjacent paragraph keeps its own sentence.",
    "Second adjacent paragraph starts independently.",
)
MIXED_LIST = "List marker preserves bold and italic words."
MIXED_TABLE = [
    ["Field", "Value"],
    ["Code label", "DOCUMATE-2026"],
    ["Status text", "Ready for review"],
]


def generate(output: Path, qa: Path, font_path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    poppler = shutil.which("pdftoppm")
    if poppler is None:
        raise RuntimeError("pdftoppm is required to generate and verify raster fixtures")
    if not font_path.is_file():
        raise FileNotFoundError(f"A Korean TrueType font is required: {font_path}")
    output.mkdir(parents=True, exist_ok=True)
    qa.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont("FixtureKorean", str(font_path)))
    pdf_path = output / "bilingual.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792), invariant=True)
    pdf.setTitle(TITLE)
    pdf.setFont("FixtureKorean", 17)
    pdf.drawString(42, 742, TITLE)
    pdf.setFont("FixtureKorean", 14)
    pdf.drawString(42, 700, HEADING)
    pdf.setFont("FixtureKorean", 11)
    for index, line in enumerate(LINES):
        pdf.drawString(42, 670 - index * 25, line)
    top, heights, widths = 506, 34, [205, 110, 205]
    for row_no, row in enumerate(TABLE):
        x = 42
        for width, value in zip(widths, row, strict=True):
            if row_no == 0:
                pdf.setFillColorRGB(0.9, 0.94, 0.97)
                pdf.rect(x, top - (row_no + 1) * heights, width, heights, fill=1, stroke=0)
            pdf.setStrokeColorRGB(0.65, 0.65, 0.65)
            pdf.rect(x, top - (row_no + 1) * heights, width, heights, fill=0, stroke=1)
            pdf.setFillColorRGB(0, 0, 0)
            pdf.drawString(x + 9, top - row_no * heights - 22, value)
            x += width
    pdf.setFont("FixtureKorean", 10)
    pdf.drawString(42, 53, "1")
    pdf.showPage()
    pdf.setFont("FixtureKorean", 17)
    pdf.drawString(42, 742, PAGE_TWO[0])
    pdf.setFont("FixtureKorean", 11)
    for index, line in enumerate(PAGE_TWO[1:]):
        pdf.drawString(42, 700 - index * 25, line)
    pdf.drawString(42, 53, "2")
    pdf.save()

    # A raster-only PDF and its identical PNG allow engine comparisons without
    # native PDF text accidentally bypassing OCR.
    subprocess.run([poppler, "-f", "1", "-l", "1", "-r", "150", "-singlefile",
                    "-png", str(pdf_path), str(output / "bilingual_scan")], check=True)
    scan = canvas.Canvas(str(output / "bilingual_scan.pdf"), pagesize=(612, 792), invariant=True)
    scan.setTitle(TITLE)
    scan.drawImage(str(output / "bilingual_scan.png"), 0, 0, width=612, height=792)
    scan.save()
    subprocess.run([poppler, "-r", "90", "-png", str(pdf_path), str(qa / "bilingual")], check=True)

    word = Document()
    section = word.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1"):
        style = word.styles[name]
        style.font.name = "Malgun Gothic"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    word.styles["Normal"].font.size = Pt(10)
    word.add_paragraph(TITLE, style="Title")
    word.add_heading(HEADING, level=1)
    for line in LINES:
        word.add_paragraph(line)
    table = word.add_table(rows=len(TABLE), cols=3)
    table.style = "Table Grid"
    for row_no, values in enumerate(TABLE):
        for cell, value in zip(table.rows[row_no].cells, values, strict=True):
            cell.text = value
            if row_no == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E5EFF7")
                cell._tc.get_or_add_tcPr().append(shade)
        if row_no == 0:
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
    word.save(output / "bilingual.docx")
    expectation = {
        "source": "Original synthetic repository fixture; generated by script/generate_document_fixtures.py",
        "title": TITLE, "heading": HEADING, "lines": LINES, "table": TABLE,
        "second_page": PAGE_TWO, "scan_dpi": 150,
        "reference_text": "\n".join([TITLE, HEADING, *LINES, *(" ".join(row) for row in TABLE)]),
    }
    (output / "bilingual_expected.json").write_text(json.dumps(expectation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def generate_mixed_formatting(output: Path, qa: Path, font_path: Path) -> None:
    """Create native-text DOCX/PDF cases without touching bilingual fixtures.

    DOCX run boundaries deliberately occur both inside words and around literal
    whitespace. PDF uses equivalent visible text, but tabs are positional and
    therefore do not have the DOCX character-preservation contract.
    """
    from datetime import datetime, timezone
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    poppler = shutil.which("pdftoppm")
    if poppler is None:
        raise RuntimeError("pdftoppm is required to verify mixed formatting fixtures")
    if not font_path.is_file():
        raise FileNotFoundError(f"A Korean TrueType font is required: {font_path}")
    output.mkdir(parents=True, exist_ok=True)
    qa.mkdir(parents=True, exist_ok=True)

    word = Document()
    word.sections[0].top_margin = word.sections[0].bottom_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        style = word.styles[name]
        style.font.name = "Malgun Gothic"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    word.styles["Normal"].font.size = Pt(10)
    word.add_paragraph(MIXED_TITLE, style="Title")
    word.add_heading(MIXED_HEADING, level=1)

    inline = word.add_paragraph()
    inline.add_run("The ")
    inline.add_run("silver").bold = True
    inline.add_run(" fox ")
    inline.add_run("follows").italic = True
    inline.add_run(" the ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), inline.part.relate_to(
        "https://example.invalid/amber-trail", RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
    ))
    hyperlink_run = OxmlElement("w:r")
    hyperlink_text = OxmlElement("w:t")
    hyperlink_text.text = "amber trail"
    hyperlink_run.append(hyperlink_text)
    hyperlink.append(hyperlink_run)
    inline._p.append(hyperlink)
    inline.add_run(" safely.")

    token = word.add_paragraph()
    token.add_run("The code is DOCU")
    token.add_run("MATE").bold = True
    token.add_run("-2026 and parse_")
    token.add_run("document").italic = True
    token.add_run(".")
    korean = word.add_paragraph()
    korean.add_run("한글")
    korean.add_run("서식경계").bold = True
    korean.add_run("를그대로")
    korean.add_run("보존").italic = True
    korean.add_run("합니다.")
    spacing = word.add_paragraph()
    spacing.add_run("Before ")
    spacing.add_run(" two spaces").bold = True
    spacing.add_run("\t")
    spacing.add_run("Tabbed segment").italic = True
    spacing.add_run("\nSoft line break stays in this paragraph.")
    for text in MIXED_ADJACENT:
        word.add_paragraph(text)
    bullet = word.add_paragraph(style="List Bullet")
    bullet.add_run("List marker preserves ")
    bullet.add_run("bold").bold = True
    bullet.add_run(" and ")
    bullet.add_run("italic").italic = True
    bullet.add_run(" words.")
    table = word.add_table(rows=len(MIXED_TABLE), cols=2)
    table.style = "Table Grid"
    for row_no, values in enumerate(MIXED_TABLE):
        for col_no, value in enumerate(values):
            cell = table.cell(row_no, col_no)
            paragraph = cell.paragraphs[0]
            # Split every value at a meaningful boundary; code has no separator.
            split = value.index(" ") if " " in value else max(1, len(value) // 2)
            paragraph.add_run(value[:split]).bold = True
            paragraph.add_run(value[split:]).italic = row_no > 0
            if row_no == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E5EFF7")
                cell._tc.get_or_add_tcPr().append(shade)
        if row_no == 0:
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)

    word.core_properties.title = MIXED_TITLE
    word.core_properties.author = "DocuMate fixtures"
    fixed_time = datetime(2026, 1, 1, tzinfo=timezone.utc)
    word.core_properties.created = word.core_properties.modified = fixed_time
    buffer = BytesIO()
    word.save(buffer)
    # Stable ZIP timestamps make repeat generation byte-for-byte reproducible.
    with ZipFile(buffer) as source, ZipFile(output / "mixed_formatting.docx", "w") as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, source.read(name))

    pdfmetrics.registerFont(TTFont("MixedFixtureKorean", str(font_path)))
    pdf_path = output / "mixed_formatting.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792), invariant=True)
    pdf.setTitle(MIXED_TITLE)
    pdf.setAuthor("DocuMate fixtures")

    def draw_runs(y: float, runs: list[tuple[str, str]], x: float = 42) -> None:
        for text, font in runs:
            pdf.setFont(font, 11)
            pdf.drawString(x, y, text)
            x += pdfmetrics.stringWidth(text, font, 11)

    pdf.setFont("Helvetica-Bold", 19)
    pdf.drawString(42, 742, MIXED_TITLE)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(42, 704, MIXED_HEADING)
    draw_runs(676, [("The ", "Helvetica"), ("silver", "Helvetica-Bold"),
                    (" fox ", "Helvetica"), ("follows", "Helvetica-Oblique"),
                    (" the amber trail safely.", "Helvetica")])
    draw_runs(650, [("The code is DOCU", "Helvetica"), ("MATE", "Helvetica-Bold"),
                    ("-2026 and parse_", "Helvetica"), ("document", "Helvetica-Oblique"),
                    (".", "Helvetica")])
    draw_runs(624, [(MIXED_KOREAN, "MixedFixtureKorean")])
    draw_runs(598, [("Before ", "Helvetica"), (" two spaces", "Helvetica-Bold")])
    draw_runs(598, [("Tabbed segment", "Helvetica-Oblique")], x=220)
    draw_runs(582, [("Soft line break stays in this paragraph.", "Helvetica")])
    for index, text in enumerate(MIXED_ADJACENT):
        draw_runs(551 - index * 26, [(text, "Helvetica")])
    draw_runs(495, [("• List marker preserves ", "Helvetica"), ("bold", "Helvetica-Bold"),
                    (" and ", "Helvetica"), ("italic", "Helvetica-Oblique"),
                    (" words.", "Helvetica")])
    top, height, widths = 465, 34, (170, 350)
    for row_no, row in enumerate(MIXED_TABLE):
        x = 42
        for width, value in zip(widths, row, strict=True):
            background = (0.9, 0.94, 0.97) if row_no == 0 else (1, 1, 1)
            pdf.setFillColorRGB(*background)
            pdf.setStrokeColorRGB(0.65, 0.65, 0.65)
            pdf.rect(x, top - (row_no + 1) * height, width, height, fill=1, stroke=1)
            pdf.setFillColorRGB(0, 0, 0)
            split = value.index(" ") if " " in value else max(1, len(value) // 2)
            draw_runs(top - row_no * height - 22, [
                (value[:split], "Helvetica-Bold"),
                (value[split:], "Helvetica-Oblique" if row_no else "Helvetica"),
            ], x=x + 9)
            x += width
    pdf.save()
    subprocess.run([poppler, "-r", "110", "-png", str(pdf_path),
                    str(qa / "mixed_formatting")], check=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("tests/fixtures/documents"))
    parser.add_argument("--qa", type=Path, default=Path("output/docling/fixture_qa"))
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/malgun.ttf"))
    parser.add_argument("--mixed-only", action="store_true",
                        help="Generate only mixed formatting DOCX/PDF fixtures; preserve bilingual fixtures")
    args = parser.parse_args()
    if args.mixed_only:
        generate_mixed_formatting(args.output, args.qa, args.font)
    else:
        generate(args.output, args.qa, args.font)
